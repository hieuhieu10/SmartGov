"""Bộ dựng DOCX hành chính theo Nghị định 30/2020/NĐ-CP.

Được tích hợp từ công cụ ``toolnd30 2/nghi-dinh-30-van-ban`` của dự án.
Luồng bước 4 dùng thể thức của DOCX dự thảo gốc và thay nội dung đã tiếp thu,
nhờ đó tệp tải xuống và tệp mở bằng trình xem DOCX là cùng một bản đúng mẫu.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Times New Roman"
_SEPARATOR = re.compile(r"^:?-{3,}:?$")


def _format_run(run, size=13, bold=False, italic=False) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _set_cell_border(cell, **kwargs) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        if edge not in kwargs:
            continue
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        for key, value in kwargs[edge].items():
            node.set(qn("w:" + key), str(value))


def _set_cell_margin(cell, top=80, start=80, bottom=80, end=80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margin = tc_pr.first_child_found_in("w:tcMar")
    if margin is None:
        margin = OxmlElement("w:tcMar")
        tc_pr.append(margin)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margin.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            margin.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths_cm: list[float]) -> None:
    table.autofit = False
    table_width = table._tbl.tblPr.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table._tbl.tblPr.append(table_width)
    table_width.set(qn("w:w"), str(sum(Cm(width).twips for width in widths_cm)))
    table_width.set(qn("w:type"), "dxa")
    for column, width in zip(table._tbl.tblGrid.gridCol_lst, widths_cm):
        column.set(qn("w:w"), str(Cm(width).twips))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(Cm(width).twips))
            cell_width.set(qn("w:type"), "dxa")


def _clear_cell(cell) -> None:
    cell._element.clear_content()


def _add_paragraph(container, text="", *, align=WD_ALIGN_PARAGRAPH.LEFT, size=13,
                   bold=False, italic=False, before=0, after=0, indent=None) -> None:
    paragraph = container.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.15
    if indent is not None:
        paragraph.paragraph_format.first_line_indent = Cm(indent)
    run = paragraph.add_run(text)
    _format_run(run, size, bold, italic)


def _clean_lines(text: str) -> list[str]:
    return [line.strip() for line in text.splitlines()
            if line.strip() and not re.fullmatch(r"[_\-\s]+", line.strip())]


def _split_table(line: str) -> list[str]:
    value = line.strip()
    if value.startswith("|"):
        value = value[1:]
    if value.endswith("|"):
        value = value[:-1]
    return [item.strip().replace(r"\|", "|") for item in value.split("|")]


def _is_table_separator(line: str) -> bool:
    cells = _split_table(line)
    return bool(cells) and all(_SEPARATOR.fullmatch(cell) for cell in cells)


def _markdown_to_content(markdown: str) -> tuple[list[str], list[dict]]:
    """Tách các đoạn và bảng Markdown để đưa vào body/phụ lục của mẫu ND30."""
    body: list[str] = []
    tables: list[dict] = []
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and _is_table_separator(lines[index + 1]):
            headers = _split_table(line)
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_split_table(lines[index]))
                index += 1
            tables.append({"headers": headers, "rows": rows})
            continue
        body.append(re.sub(r"^#{1,3}\s+", "", line))
        index += 1
    return body, tables


def _remove_repeated_formality(body: list[str], tables: list[dict], metadata: dict) -> tuple[list[str], list[dict]]:
    """Không lặp lại đầu trang khi AI trả cả thể thức trong Markdown nguồn."""
    repeated = {
        value.strip().casefold()
        for value in (
            metadata.get("co_quan_chu_quan", ""), metadata.get("co_quan_ban_hanh", ""),
            metadata.get("loai_van_ban", ""), metadata.get("trich_yeu", ""),
            "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", "Độc lập - Tự do - Hạnh phúc",
        )
        if value and value.strip()
    }
    cleaned_body = [
        value for value in body
        if value.casefold() not in repeated
        and not re.fullmatch(r"[_\-\s]+", value)
        and not value.startswith("Số:")
        and not re.match(r"^.+,\s*ngày\s+.+$", value, re.I)
    ]
    cleaned_tables = []
    for table in tables:
        all_text = " ".join(table["headers"] + [cell for row in table["rows"] for cell in row]).casefold()
        if "cộng hòa xã hội chủ nghĩa việt nam" not in all_text and "độc lập - tự do - hạnh phúc" not in all_text:
            cleaned_tables.append(table)
    return cleaned_body, cleaned_tables


def _source_metadata(source_path: str | Path, fallback_title: str) -> dict:
    """Trích thể thức có thể xác định từ DOCX gốc; không tự bịa dữ liệu thiếu."""
    metadata = {
        "co_quan_chu_quan": "", "co_quan_ban_hanh": "", "so_ky_hieu": "",
        "dia_danh": "", "ngay_ban_hanh": "", "loai_van_ban": "DỰ THẢO",
        "trich_yeu": fallback_title, "chuc_vu": "", "nguoi_ky": "", "noi_nhan": [],
    }
    try:
        source = Document(str(source_path))
    except Exception:
        return metadata
    paragraphs = [p.text.strip() for p in source.paragraphs if p.text.strip()]
    if paragraphs:
        metadata["loai_van_ban"] = paragraphs[0]
    if len(paragraphs) > 1:
        metadata["trich_yeu"] = " ".join(paragraphs[1:3])
    if source.tables:
        header = source.tables[0]
        if len(header.rows) and len(header.columns) >= 2:
            left_lines = _clean_lines(header.cell(0, 0).text)
            if left_lines:
                metadata["co_quan_chu_quan"] = " ".join(left_lines[:-1])
                metadata["co_quan_ban_hanh"] = left_lines[-1]
            header_text = header.cell(0, 1).text
            if len(header.rows) > 1:
                header_text += "\n" + header.cell(1, 1).text
                number_text = header.cell(1, 0).text
                number = re.search(r"Số:\s*([^\n]+)", number_text)
                if number:
                    metadata["so_ky_hieu"] = number.group(1).strip()
            date = re.search(r"([^\n,]+),\s*ngày\s*(.+)", header_text, re.MULTILINE)
            if date:
                metadata["dia_danh"] = date.group(1).strip()
                metadata["ngay_ban_hanh"] = date.group(2).strip()
        if len(source.tables) > 1 and len(source.tables[1].columns) >= 2:
            signature = _clean_lines(source.tables[1].cell(0, 1).text)
            if signature:
                metadata["nguoi_ky"] = signature[-1]
            if len(signature) > 1:
                metadata["chuc_vu"] = signature[-2]
            recipients = _clean_lines(source.tables[1].cell(0, 0).text)
            metadata["noi_nhan"] = [value.lstrip("- ").strip() for value in recipients[1:]]
    return metadata


def _add_body(doc: Document, items: list[str]) -> None:
    for value in items:
        if not value:
            continue
        is_heading = bool(re.match(r"^(?:[IVXLCDM]+\.|Điều\s+\d+\.?|Chương\s+[IVXLCDM]+)", value, re.I))
        if value.startswith("- "):
            _add_paragraph(doc, value[2:], after=3)
            doc.paragraphs[-1].style = "List Bullet"
        elif re.match(r"^\d+[.)]\s+", value):
            _add_paragraph(doc, re.sub(r"^\d+[.)]\s+", "", value), after=3)
            doc.paragraphs[-1].style = "List Number"
        else:
            _add_paragraph(doc, value, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=is_heading,
                           before=6 if is_heading else 0, after=5, indent=None if is_heading else 1.0)


def _add_table(doc: Document, table_data: dict) -> None:
    headers, rows = table_data["headers"], table_data["rows"]
    if not headers:
        return
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    usable_width = 25.7 if doc.sections[-1].orientation == WD_ORIENT.LANDSCAPE else 16.0
    _set_table_geometry(table, [usable_width / len(headers)] * len(headers))
    for cell, text in zip(table.rows[0].cells, headers):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_margin(cell, 90, 90, 90, 90)
        cell.text = ""
        _add_paragraph(cell, text, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, bold=True)
    for values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margin(cell, 75, 75, 75, 75)
            cell.text = ""
            _add_paragraph(cell, text, size=10)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def render_revised_draft(markdown: str, source_path: str | Path, output_path: str | Path,
                         fallback_title: str) -> None:
    """Dựng lại dự thảo bước 4 theo mẫu ND30, giữ metadata từ bản gốc."""
    data = _source_metadata(source_path, fallback_title)
    body, tables = _markdown_to_content(markdown)
    body, tables = _remove_repeated_formality(body, tables, data)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin, section.right_margin = Cm(3), Cm(2)
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(13)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)

    header = doc.add_table(rows=2, cols=2)
    _set_table_geometry(header, [5.75, 10.75])
    for cell in (cell for row in header.rows for cell in row.cells):
        _clear_cell(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_margin(cell, 0, 0, 0, 0)
        _set_cell_border(cell, top={"val": "nil"}, bottom={"val": "nil"}, left={"val": "nil"}, right={"val": "nil"})
    left, right = header.row_cells(0)
    if data["co_quan_chu_quan"]:
        _add_paragraph(left, data["co_quan_chu_quan"].upper(), align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    _add_paragraph(left, (data["co_quan_ban_hanh"] or "[TÊN CƠ QUAN, TỔ CHỨC]").upper(), align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    _add_paragraph(left, "__________", align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
    _add_paragraph(right, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    _add_paragraph(right, "Độc lập - Tự do - Hạnh phúc", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    _add_paragraph(right, "__________", align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
    number_cell, date_cell = header.row_cells(1)
    _add_paragraph(number_cell, "Số: " + (data["so_ky_hieu"] or ".../...-..."), align=WD_ALIGN_PARAGRAPH.CENTER)
    date_text = data["ngay_ban_hanh"] or "... tháng ... năm ..."
    _add_paragraph(date_cell, f"{data['dia_danh'] or '[Địa danh]'}, ngày {date_text}", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)

    _add_paragraph(doc, data["loai_van_ban"].upper(), align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True, before=8, after=2)
    if data["trich_yeu"]:
        _add_paragraph(doc, data["trich_yeu"], align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True, after=10)
    _add_body(doc, body)
    for table_data in tables:
        _add_table(doc, table_data)

    signature = doc.add_table(rows=1, cols=2)
    _set_table_geometry(signature, [8.62, 7.5])
    for cell in signature.row_cells(0):
        _clear_cell(cell)
        _set_cell_margin(cell, 0, 0, 0, 0)
        _set_cell_border(cell, top={"val": "nil"}, bottom={"val": "nil"}, left={"val": "nil"}, right={"val": "nil"})
    receive, sign = signature.row_cells(0)
    if data["noi_nhan"]:
        _add_paragraph(receive, "Nơi nhận:", size=12, bold=True, italic=True)
        for item in data["noi_nhan"]:
            _add_paragraph(receive, "- " + item, size=11)
    _add_paragraph(sign, (data["chuc_vu"] or "[CHỨC VỤ NGƯỜI KÝ]").upper(), align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True)
    _add_paragraph(sign, "(Ký, ghi rõ họ tên, đóng dấu)", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, italic=True, after=29)
    _add_paragraph(sign, data["nguoi_ky"] or "[HỌ VÀ TÊN]", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True)
    doc.save(str(output_path))
