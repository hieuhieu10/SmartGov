"""Cập nhật DOCX dự thảo tại chỗ mà không làm mất thể thức gốc."""

from __future__ import annotations

import re
import shutil
from copy import deepcopy
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

_ARTICLE = re.compile(r"^\s*Điều\s+(\d+)\s*[.:]", re.IGNORECASE)
_MARKDOWN_TABLE = re.compile(r"^\s*\|.*\|\s*$")
_INLINE_MARKDOWN = re.compile(r"(\*\*|__|`|(?<!\*)\*(?!\*)|(?<!_)_(?!_))")


@dataclass
class _ArticleBlock:
    number: str
    occurrence: int
    paragraphs: list[Paragraph] | None = None
    lines: list[str] | None = None


def _plain_markdown(line: str) -> str:
    text = re.sub(r"^\s*#{1,6}\s+", "", line.strip())
    text = _INLINE_MARKDOWN.sub("", text).replace(r"\|", "|")
    return re.sub(r"\s+", " ", text).strip()


def _markdown_lines(markdown: str) -> list[str]:
    """Lấy đoạn nội dung từ Markdown MarkItDown, bỏ bảng thể thức/tables."""
    lines: list[str] = []
    for raw in markdown.splitlines():
        if not raw.strip() or _MARKDOWN_TABLE.match(raw):
            continue
        text = _plain_markdown(raw)
        if text and not re.fullmatch(r"[_\-\s]+", text):
            lines.append(text)
    return lines


def _source_blocks(document: Document) -> list[_ArticleBlock]:
    blocks: list[_ArticleBlock] = []
    current: _ArticleBlock | None = None
    occurrences: dict[str, int] = {}
    for paragraph in document.paragraphs:
        match = _ARTICLE.match(paragraph.text)
        if match:
            number = match.group(1)
            occurrences[number] = occurrences.get(number, 0) + 1
            current = _ArticleBlock(number, occurrences[number], paragraphs=[paragraph])
            blocks.append(current)
        elif current is not None:
            current.paragraphs.append(paragraph)
    return blocks


def _revised_blocks(markdown: str) -> list[_ArticleBlock]:
    blocks: list[_ArticleBlock] = []
    current: _ArticleBlock | None = None
    occurrences: dict[str, int] = {}
    for line in _markdown_lines(markdown):
        match = _ARTICLE.match(line)
        if match:
            number = match.group(1)
            occurrences[number] = occurrences.get(number, 0) + 1
            current = _ArticleBlock(number, occurrences[number], lines=[line])
            blocks.append(current)
        elif current is not None:
            current.lines.append(line)
    return blocks


def _normalise(lines: list[str]) -> str:
    return "\n".join(re.sub(r"\s+", " ", line).strip().casefold() for line in lines if line.strip())


def _set_text_keep_format(paragraph: Paragraph, text: str) -> None:
    """Thay text nhưng giữ pPr và kiểu run đầu tiên của đoạn gốc."""
    template_rpr = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    paragraph.clear()
    run = paragraph.add_run(text)
    if template_rpr is not None:
        run._r.insert(0, template_rpr)


def _insert_after(anchor: Paragraph, template: Paragraph, text: str) -> Paragraph:
    new_element = deepcopy(template._p)
    anchor._p.addnext(new_element)
    paragraph = Paragraph(new_element, anchor._parent)
    _set_text_keep_format(paragraph, text)
    return paragraph


def _insert_before(anchor: Paragraph, template: Paragraph, text: str) -> Paragraph:
    new_element = deepcopy(template._p)
    anchor._p.addprevious(new_element)
    paragraph = Paragraph(new_element, anchor._parent)
    _set_text_keep_format(paragraph, text)
    return paragraph


def _remove_paragraph(paragraph: Paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def _replace_article(source: _ArticleBlock, revised: _ArticleBlock) -> None:
    old = source.paragraphs or []
    new = revised.lines or []
    if not old or not new:
        return
    old_keys = [_normalise([paragraph.text]) for paragraph in old]
    new_keys = [_normalise([line]) for line in new]
    # Cập nhật chỉ các đoạn khác nhau. Các đoạn equal giữ nguyên toàn bộ runs,
    # nhấn mạnh, style và định dạng đánh số của DOCX nguồn.
    opcodes = SequenceMatcher(a=old_keys, b=new_keys, autojunk=False).get_opcodes()
    for tag, old_start, old_end, new_start, new_end in reversed(opcodes):
        if tag == "equal":
            continue
        old_slice, new_slice = old[old_start:old_end], new[new_start:new_end]
        shared = min(len(old_slice), len(new_slice))
        for index in range(shared):
            _set_text_keep_format(old_slice[index], new_slice[index])
        if len(new_slice) > shared:
            template = old_slice[-1] if old_slice else (old[old_start - 1] if old_start else old[0])
            if old_slice:
                anchor = old_slice[-1]
                for text in new_slice[shared:]:
                    anchor = _insert_after(anchor, template, text)
            else:
                anchor = old[old_start] if old_start < len(old) else None
                for text in reversed(new_slice[shared:]):
                    if anchor is None:
                        anchor = _insert_after(old[-1], template, text)
                    else:
                        anchor = _insert_before(anchor, template, text)
        if len(old_slice) > shared:
            # Xóa đúng các paragraph thuộc Điều hiện tại; không động vào section/header/footer.
            for paragraph in old_slice[shared:]:
                _remove_paragraph(paragraph)


def update_revised_draft_docx(source_path: str | Path, revised_markdown: str,
                              output_path: str | Path) -> dict[str, int]:
    """Sao chép DOCX gốc và chỉ thay các Điều thay đổi theo bản AI.

    Trả về số Điều đã nhận diện/cập nhật. Không tự dựng lại tài liệu: điều này
    giữ nguyên header, footer, section, bảng, căn lề và các run định dạng gốc.
    """
    source, output = Path(source_path), Path(output_path)
    if source.suffix.casefold() != ".docx":
        raise ValueError("Chỉ có thể cập nhật giữ định dạng trực tiếp từ file DOCX gốc")
    shutil.copy2(source, output)
    document = Document(str(output))
    source_blocks = _source_blocks(document)
    revised_blocks = _revised_blocks(revised_markdown)
    if not source_blocks or not revised_blocks:
        raise ValueError("Không xác định được các Điều trong dự thảo gốc hoặc nội dung AI để cập nhật")

    revised_by_key = {(block.number, block.occurrence): block for block in revised_blocks}
    changed = 0
    matched = 0
    # Duyệt ngược để việc thêm/xóa paragraph không ảnh hưởng các block đứng trước.
    for source_block in reversed(source_blocks):
        revised_block = revised_by_key.get((source_block.number, source_block.occurrence))
        if revised_block is None:
            continue
        matched += 1
        old_lines = [paragraph.text for paragraph in source_block.paragraphs or []]
        if _normalise(old_lines) == _normalise(revised_block.lines or []):
            continue
        _replace_article(source_block, revised_block)
        changed += 1
    if not matched:
        raise ValueError("Không ghép được các Điều của bản AI với dự thảo DOCX gốc")
    document.save(str(output))
    return {"matched_articles": matched, "updated_articles": changed}
