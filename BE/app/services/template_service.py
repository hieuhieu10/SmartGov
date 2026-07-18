"""Template domain adapter and local Word exporter."""

from docx.enum.text import WD_ALIGN_PARAGRAPH

from app import database as db
from app.services.ai_client import ai_client
from app.services.word_exporter import (
    _add_content_paragraphs,
    _add_header_table,
    _add_noi_nhan,
    _add_paragraph_text,
    _add_run,
    _add_section_title,
    _add_signature_block,
    _save_doc,
    _setup_document,
)

DOC_TYPE_MAP = {
    "bao_cao": {"label": "Báo cáo", "ky_hieu": "BC"},
    "ke_hoach": {"label": "Kế hoạch", "ky_hieu": "KH"},
    "cong_van": {"label": "Công văn", "ky_hieu": "CV"},
    "quyet_dinh": {"label": "Quyết định", "ky_hieu": "QĐ"},
    "thong_bao": {"label": "Thông báo", "ky_hieu": "TB"},
    "to_trinh": {"label": "Tờ trình", "ky_hieu": "TTr"},
    "bien_ban": {"label": "Biên bản", "ky_hieu": "BB"},
    "khac": {"label": "Văn bản", "ky_hieu": "VB"},
}


class TemplateService:
    async def extract_headings(self, file_path: str, engine: str = "self_hosted") -> dict:
        return await ai_client.request(
            "/internal/template/extract",
            engine=engine,
            input_data={"stored_path": file_path},
        )

    async def generate_from_headings(
        self,
        headings: list[dict],
        doc_type_label: str,
        trich_yeu: str,
        user_input: dict,
        repo_id: str = "",
        selected_document_ids: list[str] | None = None,
        engine: str = "self_hosted",
    ) -> dict:
        selected = selected_document_ids or []
        data = await ai_client.request(
            "/internal/template/generate",
            repo_id=repo_id,
            engine=engine,
            input_data={
                "headings": headings,
                "doc_type_label": doc_type_label,
                "trich_yeu": trich_yeu,
                "user_input": user_input,
                "selected_document_ids": selected,
                "documents": await ai_client.documents(repo_id, selected or None),
            },
        )
        return data["content"]

    def build_nd30_document(
        self, doc_type: str, headings: list[dict], headings_data: dict, user_input: dict, output_path: str
    ) -> str:
        info = DOC_TYPE_MAP.get(doc_type, DOC_TYPE_MAP["khac"])
        doc = _setup_document()
        _add_header_table(
            doc,
            co_quan_chu_quan=user_input.get("co_quan_chu_quan", ""),
            ten_don_vi=user_input.get("co_quan_ban_hanh", ""),
            so_van_ban=user_input.get("so_van_ban", ""),
            ky_hieu=info["ky_hieu"],
        )
        doc.add_paragraph()
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_run(paragraph, info["label"].upper(), bold=True, size=14)
        trich_yeu = user_input.get("trich_yeu", "")
        if trich_yeu:
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_run(paragraph, trich_yeu, bold=True, size=14)
        doc.add_paragraph()
        if doc_type in ("cong_van", "to_trinh") and user_input.get("noi_nhan"):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_run(paragraph, f"Kính gửi: {user_input['noi_nhan']}", size=14)
        for heading in headings:
            key, title = heading.get("key", ""), heading.get("title", "")
            content = headings_data.get(key, "")
            if title:
                _add_section_title(doc, title)
            if content:
                _add_content_paragraphs(doc, content, title)
        closing = user_input.get("ket_luan") or f"Trên đây là {info['label']} {trich_yeu}./."
        _add_paragraph_text(doc, closing)
        _add_signature_block(
            doc,
            nguoi_ky=user_input.get("nguoi_ky", ""),
            chuc_vu=user_input.get("chuc_vu_nguoi_ky", ""),
        )
        _add_noi_nhan(doc, ["Như trên;", "Lưu: VT."])
        return _save_doc(doc, output_path)


template_service = TemplateService()
