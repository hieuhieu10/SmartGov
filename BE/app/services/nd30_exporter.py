#!/usr/bin/env python3
"""Create a Vietnamese administrative DOCX from a JSON document model."""
import argparse
import json
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Times New Roman"


def set_cell_border(cell, **kwargs):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            tag = "w:" + edge
            el = borders.find(qn(tag))
            if el is None:
                el = OxmlElement(tag)
                borders.append(el)
            for key, value in kwargs[edge].items():
                el.set(qn("w:" + key), str(value))


def set_cell_margin(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_cm):
    """Set fixed table/grid/cell widths so Word does not redistribute columns."""
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(Cm(width).twips for width in widths_cm)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for column, width in zip(grid.gridCol_lst, widths_cm):
        column.set(qn("w:w"), str(Cm(width).twips))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(Cm(width).twips))
            tc_w.set(qn("w:type"), "dxa")


def format_run(run, size=13, bold=False, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def text_para(container, text="", align=WD_ALIGN_PARAGRAPH.LEFT, size=13, bold=False,
              italic=False, before=0, after=0, indent=None, line=1.15):
    p = container.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if indent is not None:
        pf.first_line_indent = Cm(indent)
    r = p.add_run(text)
    format_run(r, size, bold, italic)
    return p


def clear_cell(cell):
    cell._element.clear_content()


def add_underline(p, width=5.2):
    p = p.add_run("\n" + "_" * int(width * 8))
    format_run(p, 10)


def normalise_list(value):
    if not value:
        return []
    return value if isinstance(value, list) else [value]


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1); run._r.append(instr); run._r.append(fld_char2)
    format_run(run, 12)


def add_body(doc, items, profile=None):
    profile = profile or {}
    for raw in normalise_list(items):
        value = str(raw).strip()
        if not value:
            continue
        if profile.get("preserve_numbering"):
            is_heading = bool(re.match(r"^[IVXLCDM]+\.\s+[A-ZÀ-Ỵ ]+$", value))
            text_para(doc, value, WD_ALIGN_PARAGRAPH.JUSTIFY, profile.get("font_size", 14),
                      bold=is_heading, before=profile.get("before", 6), after=profile.get("after", 6),
                      indent=profile.get("indent", 1.25), line=profile.get("line", 1.15))
            continue
        if re.match(r"^-\s+", value):
            p = text_para(doc, re.sub(r"^-\s+", "", value), indent=None, after=3)
            p.style = "List Bullet"
            for r in p.runs: format_run(r, 13)
        elif re.match(r"^\d+[.)]\s+", value):
            p = text_para(doc, re.sub(r"^\d+[.)]\s+", "", value), indent=None, after=3)
            p.style = "List Number"
            for r in p.runs: format_run(r, 13)
        else:
            text_para(doc, value, WD_ALIGN_PARAGRAPH.JUSTIFY, 13, after=4, indent=1.0)


def add_data_table(doc, table_data):
    """Add a simple, readable table from {headers, rows}; table geometry is explicit."""
    headers = table_data.get("headers", [])
    rows = table_data.get("rows", [])
    if not headers:
        return
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    usable_width = 25.7 if doc.sections[-1].orientation == WD_ORIENT.LANDSCAPE else 16.0
    widths = table_data.get("widths_cm") or [usable_width / len(headers)] * len(headers)
    set_table_geometry(table, widths)
    font_size = table_data.get("font_size", 10)
    for cell, text in zip(table.rows[0].cells, headers):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margin(cell, 90, 90, 90, 90)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(text)); format_run(r, font_size, True)
    for source_row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, source_row):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margin(cell, 75, 75, 75, 75)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(text)); format_run(r, font_size)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def build(data, out_path):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2)
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = FONT, Pt(13)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    layout = data.get("layout", {})
    # First block: issuing authority, national header, number and date.
    cong_van_3_cot = layout.get("cong_van_3_cot", False)
    tbl = doc.add_table(rows=3 if cong_van_3_cot else 2, cols=3 if cong_van_3_cot else 2)
    set_table_geometry(tbl, layout.get("header_widths_cm", [5.75, 10.75]))
    for c in [cell for row in tbl.rows for cell in row.cells]:
        if not cong_van_3_cot:
            clear_cell(c)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margin(c, 0, 0, 0, 0)
        set_cell_border(c, top={"val":"nil"}, bottom={"val":"nil"}, left={"val":"nil"}, right={"val":"nil"})
    if cong_van_3_cot:
        left = tbl.cell(0, 0)
        right = tbl.cell(0, 1).merge(tbl.cell(0, 2))
        date_left = tbl.cell(1, 0)
        date_right = tbl.cell(1, 1).merge(tbl.cell(1, 2))
        greeting_left = tbl.cell(2, 0)
        greeting_right = tbl.cell(2, 1).merge(tbl.cell(2, 2))
        for c in (left, right, date_left, date_right, greeting_left, greeting_right):
            clear_cell(c)
    else:
        left, right = tbl.row_cells(0)
    if data.get("do_mat"):
        text_para(left, data["do_mat"].upper(), WD_ALIGN_PARAGRAPH.CENTER, 13, True, after=3)
    if data.get("do_khan"):
        text_para(right, data["do_khan"].upper(), WD_ALIGN_PARAGRAPH.CENTER, 13, True, after=3)
    if data.get("co_quan_chu_quan"):
        text_para(left, data["co_quan_chu_quan"].upper(), WD_ALIGN_PARAGRAPH.CENTER, 13, True)
    text_para(left, data.get("co_quan_ban_hanh", "[TÊN CƠ QUAN, TỔ CHỨC]").upper(), WD_ALIGN_PARAGRAPH.CENTER, 13, True)
    text_para(left, "_" * layout.get("left_header_line_length", 10), WD_ALIGN_PARAGRAPH.CENTER, 11)
    national_size = layout.get("national_font_size", 12 if cong_van_3_cot else 13)
    text_para(right, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", WD_ALIGN_PARAGRAPH.CENTER, national_size, True)
    text_para(right, "Độc lập - Tự do - Hạnh phúc", WD_ALIGN_PARAGRAPH.CENTER, 13, True)
    text_para(right, "_" * layout.get("right_header_line_length", 10), WD_ALIGN_PARAGRAPH.CENTER, 11)

    so = data.get("so_ky_hieu") or ".../...-..."
    date = data.get("ngay_ban_hanh") or "... tháng ... năm ..."
    if layout.get("date_in_header"):
        text_para(right, f"{data.get('dia_danh', '[Địa danh]')}, ngày {date}",
                  WD_ALIGN_PARAGRAPH.CENTER, 14, italic=True, before=2.85, after=2.85)
    if not cong_van_3_cot:
        date_left, date_right = tbl.row_cells(1)
    text_para(date_left, "Số: " + so, WD_ALIGN_PARAGRAPH.CENTER, 13, before=2.85, after=2.85)
    if data.get("trich_yeu_header"):
        text_para(date_left, data["trich_yeu_header"], WD_ALIGN_PARAGRAPH.CENTER, 13, after=2.85)
    if not layout.get("date_in_header"):
        text_para(date_right, f"{data.get('dia_danh', '[Địa danh]')}, ngày {date}", WD_ALIGN_PARAGRAPH.CENTER, 14, italic=True, before=2.85, after=2.85)
    if cong_van_3_cot and data.get("kinh_gui"):
        text_para(greeting_left, "Kính gửi:", WD_ALIGN_PARAGRAPH.CENTER, 14, after=0)
        for item in normalise_list(data["kinh_gui"]):
            text_para(greeting_right, "- " + item, WD_ALIGN_PARAGRAPH.LEFT, 14, after=0)

    if data.get("hien_ten_loai", True):
        text_para(doc, data.get("loai_van_ban", "CÔNG VĂN").upper(), WD_ALIGN_PARAGRAPH.CENTER, 14, True,
                  before=layout.get("title_before", 0), after=2)
    trich_yeu = data.get("trich_yeu")
    title_lines = normalise_list(data.get("trich_yeu_dong")) or normalise_list(trich_yeu)
    for title_line in title_lines:
        text_para(doc, title_line, WD_ALIGN_PARAGRAPH.CENTER, 14, True, after=0)
    if title_lines and data.get("hien_ten_loai", True):
        if layout.get("title_separator"):
            text_para(doc, "___________", WD_ALIGN_PARAGRAPH.CENTER, 12, after=16)
        else:
            doc.paragraphs[-1].paragraph_format.space_after = Pt(9)
    if data.get("kinh_gui") and not cong_van_3_cot:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.15
        r = p.add_run("Kính gửi: "); format_run(r, 13, True)
        r = p.add_run("; ".join(normalise_list(data["kinh_gui"]))); format_run(r, 13)

    add_body(doc, data.get("noi_dung", []), data.get("body_profile"))

    sig = doc.add_table(rows=1, cols=2)
    set_table_geometry(sig, layout.get("signature_widths_cm", [8.62, 7.5]))
    for c in sig.row_cells(0):
        clear_cell(c)
        set_cell_margin(c, 0, 0, 0, 0)
        set_cell_border(c, top={"val":"nil"}, bottom={"val":"nil"}, left={"val":"nil"}, right={"val":"nil"})
    left, right = sig.row_cells(0)
    recipients = normalise_list(data.get("noi_nhan"))
    if recipients:
        p = text_para(left, "Nơi nhận:", size=12, bold=True, italic=True)
        for item in recipients:
            text_para(left, "- " + item, size=11.5, after=0)
    for line in normalise_list(data.get("chu_ky_dong")) or normalise_list(data.get("thua_lenh")):
        text_para(right, line.upper(), WD_ALIGN_PARAGRAPH.CENTER, 14, True)
    text_para(right, data.get("chuc_vu", "[CHỨC VỤ NGƯỜI KÝ]").upper(), WD_ALIGN_PARAGRAPH.CENTER, 14, True)
    if not data.get("chu_ky_dong"):
        text_para(right, "(Ký, ghi rõ họ tên, đóng dấu)", WD_ALIGN_PARAGRAPH.CENTER, 12, italic=True, after=29)
    else:
        text_para(right, "", WD_ALIGN_PARAGRAPH.CENTER, 14, after=29)
    text_para(right, data.get("nguoi_ky", "[HỌ VÀ TÊN]"), WD_ALIGN_PARAGRAPH.CENTER, 14, True)
    if data.get("nguoi_soan_thao"):
        text_para(left, data["nguoi_soan_thao"], size=9, after=0)

    for appendix in data.get("phu_luc", []):
        if appendix.get("landscape") and doc.sections[-1].orientation != WD_ORIENT.LANDSCAPE:
            app_sec = doc.add_section(WD_SECTION.NEW_PAGE)
            app_sec.orientation = WD_ORIENT.LANDSCAPE
            app_sec.page_width, app_sec.page_height = Cm(29.7), Cm(21)
            app_sec.top_margin, app_sec.bottom_margin = Cm(1.5), Cm(1.5)
            app_sec.left_margin, app_sec.right_margin = Cm(2), Cm(2)
        else:
            doc.add_page_break()
        text_para(doc, "PHỤ LỤC", WD_ALIGN_PARAGRAPH.CENTER, 14, True)
        text_para(doc, appendix.get("tieu_de", ""), WD_ALIGN_PARAGRAPH.CENTER, 13, True, after=10)
        add_body(doc, appendix.get("noi_dung", []), appendix.get("body_profile"))
        for table_data in appendix.get("bang", []):
            add_data_table(doc, table_data)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer)
    # OOXML core-property titles are limited to 255 characters; this never affects visible text.
    doc.core_properties.title = str(data.get("trich_yeu", data.get("loai_van_ban", "Văn bản hành chính")))[:255]
    doc.save(out_path)


def render_docx(docx_path, render_dir):
    """Render all pages via LibreOffice, with macOS Quick Look thumbnail fallback."""
    docx_path, render_dir = Path(docx_path).resolve(), Path(render_dir).resolve()
    if not docx_path.exists():
        raise FileNotFoundError(f"Không tìm thấy DOCX: {docx_path}")
    render_dir.mkdir(parents=True, exist_ok=True)
    converter = shutil.which("soffice") or shutil.which("libreoffice")
    poppler = shutil.which("pdftoppm")
    if converter and poppler:
        with tempfile.TemporaryDirectory(prefix="nd30-render-") as temp:
            temp_path = Path(temp)
            profile = temp_path / "lo-profile"
            command = [converter, "--headless", f"-env:UserInstallation=file://{profile}",
                       "--convert-to", "pdf", "--outdir", str(temp_path), str(docx_path)]
            subprocess.run(command, check=True, capture_output=True, text=True)
            pdf = temp_path / (docx_path.stem + ".pdf")
            if not pdf.exists():
                raise RuntimeError("LibreOffice không tạo được PDF.")
            final_pdf = render_dir / (docx_path.stem + ".pdf")
            shutil.copy2(pdf, final_pdf)
            subprocess.run([poppler, "-png", "-r", "160", str(final_pdf),
                            str(render_dir / "page")], check=True)
        pages = sorted(render_dir.glob("page-*.png"))
        print(f"Đã render {len(pages)} trang: {render_dir}")
        return
    # Quick Look can at least produce a usable first-page visual check on macOS.
    qlmanage = shutil.which("qlmanage")
    if qlmanage:
        with tempfile.TemporaryDirectory(prefix="nd30-quicklook-") as temp:
            subprocess.run([qlmanage, "-t", "-s", "2400", "-o", temp, str(docx_path)],
                           check=True, capture_output=True, text=True)
            thumbnail = Path(temp) / (docx_path.name + ".png")
            if thumbnail.exists():
                shutil.copy2(thumbnail, render_dir / "page-1.png")
                print("Đã tạo page-1.png bằng Quick Look. Đây chỉ là ảnh trang đầu; "
                      "cài LibreOffice + Poppler để render và QA toàn bộ trang.")
                return
    raise RuntimeError("Không có LibreOffice+Poppler hoặc Quick Look để render DOCX.")


def main():
    parser = argparse.ArgumentParser(description="Tạo DOCX theo thể thức Nghị định 30/2020/NĐ-CP")
    src = parser.add_mutually_exclusive_group(required=True)
    src.add_argument("--input", help="Tệp JSON UTF-8")
    src.add_argument("--text", help="Tệp văn bản thuần cho phần nội dung")
    src.add_argument("--render-docx", help="Chỉ render DOCX đã có, không tạo văn bản mới")
    parser.add_argument("--output", help="Tệp .docx đầu ra")
    parser.add_argument("--render", action="store_true", help="Render DOCX sau khi tạo")
    parser.add_argument("--render-dir", help="Thư mục PDF/PNG render; mặc định <tên>-render")
    parser.add_argument("--co-quan-ban-hanh")
    parser.add_argument("--dia-danh")
    parser.add_argument("--nguoi-ky")
    parser.add_argument("--chuc-vu")
    parser.add_argument("--loai-van-ban", default="CÔNG VĂN")
    args = parser.parse_args()
    if args.render_docx:
        render_docx(args.render_docx, args.render_dir or (str(Path(args.render_docx).with_suffix("")) + "-render"))
        return
    if not args.output:
        parser.error("--output là bắt buộc khi tạo văn bản.")
    if args.input:
        data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    else:
        body = Path(args.text).read_text(encoding="utf-8")
        data = {"noi_dung": [p.strip() for p in body.split("\n\n") if p.strip()],
                "co_quan_ban_hanh": args.co_quan_ban_hanh, "dia_danh": args.dia_danh,
                "nguoi_ky": args.nguoi_ky, "chuc_vu": args.chuc_vu,
                "loai_van_ban": args.loai_van_ban}
    build(data, args.output)
    print(f"Đã tạo: {args.output}")
    if args.render:
        render_docx(args.output, args.render_dir or (str(Path(args.output).with_suffix("")) + "-render"))


if __name__ == "__main__":
    main()
