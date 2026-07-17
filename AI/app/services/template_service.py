"""
Template Service — Extract headings from uploaded files, generate NĐ 30 documents.

Dual-engine support:
  - NotebookLM: Upload to temp notebook → extract/generate via chat API
  - Self-hosted: Convert to markdown → extract via vLLM → generate with RAG context

Flow:
1. User uploads a file (docx/pdf/img)
2. AI extracts main headings → JSON in DB
3. User selects template + repository → AI generates content per heading
4. word_exporter builds NĐ 30 formatted document
"""

import json
import logging
import re
import asyncio
from datetime import datetime
from pathlib import Path
from typing import Optional

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.config import settings

logger = logging.getLogger(__name__)


# ─── Prompts ─────────────────────────────────────────────────────────

EXTRACT_HEADINGS_PROMPT = """Bạn là chuyên viên phân tích cấu trúc văn bản hành chính Việt Nam.

Hãy phân tích file tài liệu mẫu đã upload và trích xuất cấu trúc đầu mục chính.

Trả về JSON với cấu trúc CHÍNH XÁC:

{
    "doc_type": "Mã loại văn bản (một trong: bao_cao, ke_hoach, cong_van, quyet_dinh, thong_bao, to_trinh, bien_ban, khac)",
    "doc_type_label": "Tên loại văn bản tiếng Việt (VD: Báo cáo, Kế hoạch, Công văn...)",
    "trich_yeu": "Trích yếu/tiêu đề nội dung chính của văn bản mẫu này",
    "headings": [
        {
            "key": "ten_viet_khong_dau_gach_duoi",
            "title": "Tiêu đề đầu mục gốc (VD: I. CÔNG TÁC CHỈ ĐẠO, ĐIỀU HÀNH)",
            "description": "Mô tả ngắn gọn nội dung cần viết cho đầu mục này",
            "required": true
        }
    ]
}

QUY TẮC:
1. CHỈ trả về JSON thuần, KHÔNG markdown, KHÔNG giải thích
2. Mỗi đầu mục lớn (I, II, III, IV, V...) là MỘT entry trong headings
3. KHÔNG tách nhỏ nội dung trong cùng mục lớn
4. key phải viết thường, không dấu, dùng gạch dưới (VD: cong_tac_chi_dao, ket_qua_thuc_hien)
5. Mỗi key phải DUY NHẤT, không trùng lặp
6. doc_type phải là một trong: bao_cao, ke_hoach, cong_van, quyet_dinh, thong_bao, to_trinh, bien_ban, khac
7. Nếu không rõ loại VB, dùng "khac"
8. Chỉ trích xuất đầu mục CẤP LỚN (I, II, III...), KHÔNG liệt kê mục con (1, 2, 3, a, b, c)
9. Tổng số headings nên từ 3-10 mục
"""


GENERATE_HEADING_CONTENT_PROMPT = """Bạn là chuyên viên soạn thảo văn bản hành chính nhà nước Việt Nam bậc cao.

Nhiệm vụ: Dựa vào tài liệu tham khảo, hãy viết nội dung cho phần sau của văn bản {doc_type_label}:

ĐẦU MỤC: {heading_title}
MÔ TẢ: {heading_description}
TRÍCH YẾU VĂN BẢN: {trich_yeu}

{user_hint_section}

{context_section}

HƯỚNG DẪN:
1. PHẢI trích xuất và sử dụng thông tin THỰC TẾ từ tài liệu tham khảo
2. Văn phong hành chính, trang trọng, ngắn gọn nhưng đầy đủ ý
3. Trình bày theo đúng thể thức NĐ 30/2020/NĐ-CP
4. Dùng đánh số (1., 2., 3.) cho các mục con, gạch đầu dòng (-) cho chi tiết
5. Nếu không có thông tin liên quan trong tài liệu, viết nội dung khung phù hợp
6. Dùng \\n để xuống dòng giữa các mục/đoạn
7. KHÔNG viết lại tiêu đề đầu mục — hệ thống tự thêm
8. CHỈ trả về nội dung thuần (plain text), KHÔNG markdown, KHÔNG JSON
"""


# ─── Document Type Mapping ───────────────────────────────────────────

DOC_TYPE_MAP = {
    "bao_cao": {"label": "Báo cáo", "ky_hieu": "BC"},
    "ke_hoach": {"label": "Kế hoạch", "ky_hieu": "KH"},
    "cong_van": {"label": "Công văn", "ky_hieu": "CV"},
    "quyet_dinh": {"label": "Quyết định", "ky_hieu": "QĐ"},
    "thong_bao": {"label": "Thông báo", "ky_hieu": "TB"},
    "to_trinh": {"label": "Tờ trình", "ky_hieu": "TTr"},
    "bien_ban": {"label": "Biên bản", "ky_hieu": "BB"},
    "khac": {"label": "Văn bản", "ky_hieu": "VB"},
}


class TemplateService:
    """Extract headings from templates and generate NĐ 30 documents."""

    @staticmethod
    def _clean_content(text: str) -> str:
        """Clean AI-generated content: fix literal \\n, excess blank lines, etc."""
        import re
        # Replace literal \n with actual newline
        text = text.replace('\\n', '\n')
        # Remove \r
        text = text.replace('\r', '')
        # Collapse 3+ consecutive newlines into 2
        text = re.sub(r'\n{3,}', '\n\n', text)
        # Remove leading/trailing whitespace per line
        lines = [line.strip() for line in text.split('\n')]
        # Remove empty lines at start/end
        while lines and not lines[0]:
            lines.pop(0)
        while lines and not lines[-1]:
            lines.pop()
        return '\n'.join(lines)

    # ═══════════════════════════════════════════════════════════════════
    # Step 1: Extract Headings from uploaded file
    # ═══════════════════════════════════════════════════════════════════

    async def extract_headings(self, file_path: str) -> dict:
        """
        Extract main headings from an uploaded file.

        Routes to NotebookLM or self-hosted engine based on config.

        Returns: {
            doc_type: str,
            doc_type_label: str,
            trich_yeu: str,
            headings: [{key, title, description, required}]
        }
        """
        if settings.is_self_hosted:
            return await self._extract_headings_self_hosted(file_path)
        else:
            return await self._extract_headings_notebooklm(file_path)

    async def _extract_headings_notebooklm(self, file_path: str) -> dict:
        """Extract headings using NotebookLM temp notebook."""
        from app.services.notebooklm_service import notebooklm_service

        logger.info(f"[NotebookLM] Extracting headings from: {file_path}")

        notebook_id = None
        try:
            notebook_id = await notebooklm_service.create_notebook(
                f"Template: {Path(file_path).stem}"
            )
            logger.info(f"Created temp notebook {notebook_id}")

            await notebooklm_service.upload_source(notebook_id, file_path)
            logger.info(f"Uploaded source to notebook {notebook_id}")

            # Wait for NotebookLM to process
            await asyncio.sleep(5)

            raw = await notebooklm_service.chat_ask(notebook_id, EXTRACT_HEADINGS_PROMPT)
            logger.info(f"Extraction response: {len(raw)} chars")

            return self._parse_heading_result(raw)

        finally:
            if notebook_id:
                try:
                    await notebooklm_service.delete_notebook(notebook_id)
                    logger.info(f"Deleted temp notebook {notebook_id}")
                except Exception as e:
                    logger.warning(f"Could not delete temp notebook: {e}")

    async def _extract_headings_self_hosted(self, file_path: str) -> dict:
        """
        Extract headings using self-hosted engine:
        1. Convert file to markdown via ETL pipeline
        2. Send markdown content to vLLM for heading extraction
        """
        from app.services.document_converter import document_converter
        from app.services.llm_service import llm_service

        logger.info(f"[Self-Hosted] Extracting headings from: {file_path}")

        # Step 1: Convert to markdown
        markdown_text = document_converter.convert_to_markdown(file_path)
        if not markdown_text.strip():
            raise ValueError("Không thể đọc nội dung file. Vui lòng thử file khác.")

        # Step 2: Truncate if too long (vLLM context limit)
        max_content = 12000
        if len(markdown_text) > max_content:
            markdown_text = markdown_text[:max_content] + "\n... (nội dung đã cắt)"

        # Step 3: Ask vLLM
        user_prompt = f"""NỘI DUNG FILE MẪU:
---
{markdown_text}
---

{EXTRACT_HEADINGS_PROMPT}"""

        raw = await llm_service.chat(
            system_prompt="Bạn là chuyên viên phân tích cấu trúc văn bản hành chính Việt Nam.",
            user_prompt=user_prompt,
            temperature=0.1,
            max_tokens=4096,
        )

        return self._parse_heading_result(raw)

    def _parse_heading_result(self, raw: str) -> dict:
        """Parse heading extraction result from AI response."""
        from app.services.notebooklm_service import notebooklm_service

        json_str = notebooklm_service._extract_json(raw)
        if json_str:
            result = json.loads(json_str)
        else:
            raise ValueError("Hệ thống không trả về kết quả hợp lệ. Vui lòng thử lại.")

        # Validate
        headings = result.get("headings", [])
        if not headings:
            raise ValueError("Không trích xuất được đầu mục nào từ file")

        # Ensure valid doc_type
        doc_type = result.get("doc_type", "khac")
        if doc_type not in DOC_TYPE_MAP:
            doc_type = "khac"
        result["doc_type"] = doc_type
        result["doc_type_label"] = DOC_TYPE_MAP[doc_type]["label"]

        # Deduplicate keys
        seen_keys = set()
        unique_headings = []
        for h in headings:
            key = h.get("key", "")
            if key and key not in seen_keys:
                seen_keys.add(key)
                unique_headings.append(h)
        result["headings"] = unique_headings

        logger.info(
            f"Extracted {len(unique_headings)} headings, "
            f"doc_type={doc_type}, trich_yeu={result.get('trich_yeu', '')[:50]}"
        )
        return result

    # ═══════════════════════════════════════════════════════════════════
    # Step 2: Generate content for each heading
    # ═══════════════════════════════════════════════════════════════════

    async def generate_from_headings(
        self,
        headings: list[dict],
        notebook_id: str,
        doc_type_label: str,
        trich_yeu: str,
        user_input: dict,
        repo_id: str = "",
        selected_document_ids: list[str] | None = None,
    ) -> dict:
        """
        Generate content for each heading.

        Routes to NotebookLM or self-hosted engine based on config.
        """
        if settings.is_self_hosted:
            return await self._generate_self_hosted(
                headings, doc_type_label, trich_yeu, user_input, repo_id,
                selected_document_ids=selected_document_ids or [],
            )
        else:
            source_filter = await self._build_source_filter(selected_document_ids or [])
            return await self._generate_notebooklm(
                headings, notebook_id, doc_type_label, trich_yeu, user_input,
                source_filter=source_filter,
            )

    async def _generate_notebooklm(
        self,
        headings: list[dict],
        notebook_id: str,
        doc_type_label: str,
        trich_yeu: str,
        user_input: dict,
        source_filter: str = "",
    ) -> dict:
        """Generate heading content using NotebookLM."""
        from app.services.notebooklm_service import notebooklm_service

        result = {}
        for heading in headings:
            key = heading.get("key", "")
            title = heading.get("title", "")
            description = heading.get("description", "")

            if not key:
                continue

            # Check if user provided content for this heading
            user_val = user_input.get(key, "")
            if user_val and len(user_val.strip()) > 30:
                result[key] = user_val.strip()
                logger.info(f"Heading '{key}': using user-provided content")
                continue

            user_hint_section = ""
            if user_val:
                user_hint_section = f"GỢI Ý TỪ NGƯỜI DÙNG: {user_val}"
            if source_filter:
                user_hint_section = f"{source_filter}\n\n{user_hint_section}".strip()

            prompt = GENERATE_HEADING_CONTENT_PROMPT.format(
                doc_type_label=doc_type_label,
                heading_title=title,
                heading_description=description,
                trich_yeu=trich_yeu,
                user_hint_section=user_hint_section,
                context_section="",
            )

            try:
                content = await notebooklm_service.chat_ask(notebook_id, prompt)
                content = content.strip()
                if content.startswith(title):
                    content = content[len(title):].lstrip(".:- ").strip()
                content = self._clean_content(content)
                result[key] = content
                logger.info(f"Heading '{key}': generated {len(content)} chars")
            except Exception as e:
                logger.error(f"Failed to generate content for heading '{key}': {e}")
                result[key] = f"[Nội dung chưa sinh được cho mục: {title}]"

        return result

    async def _build_source_filter(self, selected_document_ids: list[str]) -> str:
        if not selected_document_ids:
            return ""
        from app.services.document_scanner import get_request_documents
        selected = set(selected_document_ids[:30])
        names = [
            doc.get("filename", "")
            for doc in get_request_documents()
            if str(doc.get("id")) in selected and doc.get("filename")
        ]
        if not names:
            return ""
        lines = "\n".join(f"- {name}" for name in names)
        return (
            "GIỚI HẠN TÀI LIỆU NGUỒN: Người dùng đã chọn tài liệu cụ thể. "
            "Chỉ sử dụng thông tin từ các tài liệu sau; bỏ qua tài liệu khác trong kho/sổ ghi chú:\n"
            f"{lines}"
        )

    async def _generate_self_hosted(
        self,
        headings: list[dict],
        doc_type_label: str,
        trich_yeu: str,
        user_input: dict,
        repo_id: str,
        selected_document_ids: list[str] | None = None,
    ) -> dict:
        """
        Generate heading content using self-hosted engine:
        1. Scan documents for relevant content via DocumentScanner
        2. Generate content via vLLM
        """
        from app.services.document_scanner import document_scanner
        from app.services.llm_service import llm_service

        # Step 1: Scan repository for relevant content
        outline = f"{trich_yeu}\n" + "\n".join(
            f"- {h.get('title', '')}: {h.get('description', '')}" for h in headings
        )
        scan_results = await document_scanner.scan_repository(
            repo_id=repo_id,
            outline=outline,
            document_ids=selected_document_ids or [],
        )
        context = document_scanner.format_scanner_results(scan_results)

        result = {}
        for heading in headings:
            key = heading.get("key", "")
            title = heading.get("title", "")
            description = heading.get("description", "")

            if not key:
                continue

            # Check if user provided content for this heading
            user_val = user_input.get(key, "")
            if user_val and len(user_val.strip()) > 30:
                result[key] = user_val.strip()
                logger.info(f"Heading '{key}': using user-provided content")
                continue

            # Build prompt
            user_hint_section = ""
            if user_val:
                user_hint_section = f"GỢI Ý TỪ NGƯỜI DÙNG: {user_val}"

            context_section = f"DỮ LIỆU THAM KHẢO TỪ KHO:\n{context}" if context else ""

            prompt = GENERATE_HEADING_CONTENT_PROMPT.format(
                doc_type_label=doc_type_label,
                heading_title=title,
                heading_description=description,
                trich_yeu=trich_yeu,
                user_hint_section=user_hint_section,
                context_section=context_section,
            )

            # Generate via vLLM
            try:
                content = await llm_service.chat(
                    system_prompt="Bạn là chuyên viên soạn thảo văn bản hành chính nhà nước Việt Nam.",
                    user_prompt=prompt,
                    temperature=0.2,
                    max_tokens=4096,
                )
                content = content.strip()
                if content.startswith(title):
                    content = content[len(title):].lstrip(".:- ").strip()
                content = self._clean_content(content)
                result[key] = content
                logger.info(
                    f"[Self-Hosted] Heading '{key}': {len(content)} chars"
                )
            except Exception as e:
                logger.error(f"Failed to generate content for heading '{key}': {e}")
                result[key] = f"[Nội dung chưa sinh được cho mục: {title}]"

        return result

    # ═══════════════════════════════════════════════════════════════════
    # Step 3: Build NĐ 30 formatted Word document
    # ═══════════════════════════════════════════════════════════════════

    def build_nd30_document(
        self,
        doc_type: str,
        headings: list[dict],
        headings_data: dict,
        user_input: dict,
        output_path: str,
    ) -> str:
        """
        Build a Word document following NĐ 30/2020/NĐ-CP formatting.

        Uses the same formatting utilities as word_exporter.py:
        - A4, Times New Roman 14pt, margins 3-2-2-2 cm
        - Header table: cơ quan | quốc hiệu
        - Section titles (bold)
        - Content paragraphs (justified, first-line indent)

        Args:
            doc_type: one of DOC_TYPE_MAP keys
            headings: [{key, title, description, required}]
            headings_data: {key: content_text} from generate_from_headings
            user_input: {key: user_text, co_quan_ban_hanh, trich_yeu, etc.}
            output_path: where to save the .docx
        """
        from app.services.word_exporter import (
            _setup_document, _add_header_table, _add_section_title,
            _add_content_paragraphs, _add_signature_block, _add_noi_nhan,
            _save_doc, _add_paragraph_text, _add_run,
        )

        doc_info = DOC_TYPE_MAP.get(doc_type, DOC_TYPE_MAP["khac"])

        doc = _setup_document()

        # ── Header Table ──
        _add_header_table(
            doc,
            co_quan_chu_quan=user_input.get("co_quan_chu_quan", ""),
            ten_don_vi=user_input.get("co_quan_ban_hanh", ""),
            so_van_ban=user_input.get("so_van_ban", ""),
            ky_hieu=doc_info["ky_hieu"],
        )

        # ── Document Title ──
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_run(p, doc_info["label"].upper(), bold=True, size=14)

        trich_yeu = user_input.get("trich_yeu", "")
        if trich_yeu:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_run(p, trich_yeu, bold=True, size=14)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_run(p, "________", size=14)

        doc.add_paragraph()

        # ── Kính gửi (for cong_van, to_trinh) ──
        if doc_type in ("cong_van", "to_trinh"):
            noi_nhan = user_input.get("noi_nhan", "")
            if noi_nhan:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _add_run(p, f"Kính gửi: {noi_nhan}", size=14)
                doc.add_paragraph()

        # ── Dynamic Sections from Headings ──
        for heading in headings:
            key = heading.get("key", "")
            title = heading.get("title", "")
            content = headings_data.get(key, "")

            if not content and not title:
                continue

            # Add section title (bold, like "I. CÔNG TÁC CHỈ ĐẠO")
            _add_section_title(doc, title)

            # Add content paragraphs
            if content:
                _add_content_paragraphs(doc, content, title)

        # ── Closing paragraph ──
        closing = user_input.get("ket_luan", "")
        if closing:
            _add_paragraph_text(doc, closing)
        else:
            co_quan = user_input.get("co_quan_ban_hanh", "")
            if doc_type == "bao_cao":
                _add_paragraph_text(doc, f"Trên đây là {doc_info['label']} {trich_yeu}. {co_quan}./.")
            elif doc_type == "ke_hoach":
                _add_paragraph_text(doc, f"Trên đây là {doc_info['label']} {trich_yeu}. Trong quá trình thực hiện, nếu có vướng mắc, phát sinh, các đơn vị phản ánh kịp thời về {co_quan} để xem xét, giải quyết./.")
            elif doc_type == "to_trinh":
                noi_nhan = user_input.get("noi_nhan", "")
                _add_paragraph_text(doc, f"{co_quan} kính trình {noi_nhan} xem xét, quyết định./.")
            elif doc_type == "thong_bao":
                _add_paragraph_text(doc, f"{co_quan} thông báo để các cơ quan, đơn vị, cá nhân liên quan biết và thực hiện./.")
            elif doc_type == "cong_van":
                _add_paragraph_text(doc, f"{co_quan} trân trọng kính gửi./.")
            else:
                _add_paragraph_text(doc, f"Trên đây là {doc_info['label']} {trich_yeu}./.")

        # ── Signature Block ──
        _add_signature_block(
            doc,
            nguoi_ky=user_input.get("nguoi_ky", ""),
            chuc_vu=user_input.get("chuc_vu_nguoi_ky", ""),
        )

        # ── Nơi nhận ──
        noi_nhan_list = ["Như trên;", "Lưu: VT."]
        if doc_type == "ke_hoach":
            noi_nhan_list = ["Như trên;", "Các phòng, ban, đơn vị thuộc Sở;", "Lưu: VT."]
        _add_noi_nhan(doc, noi_nhan_list)

        return _save_doc(doc, output_path)


# Singleton
template_service = TemplateService()
